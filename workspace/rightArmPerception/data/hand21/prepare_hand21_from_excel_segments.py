#!/usr/bin/env python3
"""Prepare per-video hand21 gesture clips from MediaPipe annotations.

This script focuses on *preparation only*:
1. Read gesture segments from an Excel/CSV annotation file.
2. Group segments by source video.
3. Optionally crop each source video into per-segment clip videos.
4. Run MediaPipe extraction for each source video or each cropped clip.
5. Convert the extracted tracks into hand21 JSON clips organized by source
   video under ``--output-root``.
6. Write per-video manifests plus a global ``category_annotation.json``.

Train/val/test splitting is intentionally left to a separate script so this
preparation stage can be incrementally updated per source video.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import os
import random
import shutil
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
from collections import OrderedDict
from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np


def find_repo_root(start: Path) -> Path:
    current = start.resolve()
    for candidate in (current, *current.parents):
        if (candidate / "data").is_dir() and (candidate / "training").is_dir():
            return candidate
    raise RuntimeError(f"Cannot locate repo root from {start}")


REPO_ROOT = find_repo_root(Path(__file__).parent)
EXTRACTOR_PATH = REPO_ROOT / "data/gesture_data/extract_hand_roi_with_kpts.py"


def log_status(message: str) -> None:
    timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] {message}", flush=True)


def progress_interval(total: int) -> int:
    if total <= 20:
        return 1
    return max(1, min(500, total // 20))


def should_report_progress(current: int, total: int) -> bool:
    if total <= 0:
        return True
    if current <= 1 or current == total:
        return True
    return current % progress_interval(total) == 0


def format_progress(current: int, total: int) -> str:
    if total <= 0:
        return f"{current}/0"
    return f"{current}/{total} ({current / total * 100:.1f}%)"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Prepare per-source-video hand21 gesture clips from Excel/CSV "
            "annotations without performing train/val/test splitting."
        )
    )
    parser.add_argument("--video-root", type=Path, required=True, help="Root directory containing source mp4 files.")
    parser.add_argument("--annotation-file", type=Path, required=True, help="Excel/CSV file with segment annotations.")
    parser.add_argument(
        "--output-root",
        type=Path,
        required=True,
        help="Output dataset root organized by source video.",
    )
    parser.add_argument(
        "--kpts-root",
        type=Path,
        default=None,
        help="Directory for MediaPipe annotations. Defaults to <output-root>_video_kpts.",
    )
    parser.add_argument(
        "--crop-first",
        action="store_true",
        help=(
            "Crop annotated video segments into clip videos before MediaPipe extraction. "
            "When enabled, extraction runs on the cropped clips instead of full source videos."
        ),
    )
    parser.add_argument(
        "--crop-root",
        type=Path,
        default=None,
        help="Directory for cropped clip videos when --crop-first is enabled. Defaults to <output-root>_segments_videos.",
    )
    parser.add_argument("--sheet-name", type=str, default=None, help="Excel sheet name. Uses the first sheet by default.")
    parser.add_argument("--video-col", type=str, default="video_name", help="Column name for video path/name.")
    parser.add_argument("--label-col", type=str, default="label", help="Column name for gesture label.")
    parser.add_argument("--start-col", type=str, default="start_frame", help="Column name for start frame.")
    parser.add_argument("--end-col", type=str, default="end_frame", help="Column name for end frame.")
    parser.add_argument(
        "--annotation-format",
        choices=["auto", "table", "compact_4col"],
        default="auto",
        help=(
            "Annotation layout. 'table' expects a header row. "
            "'compact_4col' expects columns [video(optional), start, end, label] "
            "with video inherited from the last non-empty cell. "
            "'auto' detects between the two."
        ),
    )
    parser.add_argument(
        "--split-col",
        type=str,
        default=None,
        help="Deprecated for this script. If provided, the value is preserved in metadata only.",
    )
    parser.add_argument(
        "--frame-base",
        type=int,
        choices=[0, 1],
        default=0,
        help="Whether start/end frame in the annotation file are 0-based or 1-based.",
    )
    parser.add_argument("--pre-pad", type=int, default=0, help="Frames to extend before the annotated segment.")
    parser.add_argument("--post-pad", type=int, default=0, help="Frames to extend after the annotated segment.")
    parser.add_argument("--min-frames", type=int, default=1, help="Minimum clip length to keep after padding/clamping.")
    parser.add_argument(
        "--min-valid-frames",
        type=int,
        default=0,
        help="Minimum number of valid frames required for a clip to be kept.",
    )
    parser.add_argument(
        "--min-valid-ratio",
        type=float,
        default=0.0,
        help="Minimum valid-frame ratio required for a clip to be kept.",
    )
    parser.add_argument(
        "--val-ratio",
        type=float,
        default=0.2,
        help="Deprecated for this script. Kept for backward compatibility and ignored.",
    )
    parser.add_argument(
        "--test-ratio",
        type=float,
        default=0.0,
        help="Deprecated for this script. Kept for backward compatibility and ignored.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=42,
        help="Deprecated for this script. Kept for backward compatibility and ignored.",
    )
    parser.add_argument(
        "--handedness",
        type=str,
        choices=["Left", "Right", "any"],
        default="any",
        help="Preferred handedness when auto-selecting a track from a video.",
    )
    parser.add_argument(
        "--track-id",
        type=int,
        default=None,
        help="Optional explicit track ID to use for every video. Usually keep unset.",
    )
    parser.add_argument(
        "--use-smooth",
        dest="use_smooth",
        action="store_true",
        help="Use kpts_xy_smooth when available. Default: enabled.",
    )
    parser.add_argument(
        "--no-use-smooth",
        dest="use_smooth",
        action="store_false",
        help="Disable smoothed keypoints and use raw kpts_xy.",
    )
    parser.set_defaults(use_smooth=True)
    parser.add_argument(
        "--joint-score-source",
        choices=["hand_score", "ones"],
        default="hand_score",
        help="How to fill the per-joint score in generated clip JSON.",
    )
    parser.add_argument(
        "--skip-extract",
        action="store_true",
        help="Skip running MediaPipe extraction and reuse existing --kpts-root annotations.",
    )
    parser.add_argument("--model-path", type=str, default="hand_landmarker.task", help="MediaPipe model path.")
    parser.add_argument("--num-hands", type=int, default=2, help="Maximum hands per frame for MediaPipe extraction.")
    parser.add_argument("--padding-ratio", type=float, default=0.25, help="Padding ratio forwarded to extractor.")
    parser.add_argument("--min-hand-detection-confidence", type=float, default=0.5)
    parser.add_argument("--min-hand-presence-confidence", type=float, default=0.5)
    parser.add_argument("--min-tracking-confidence", type=float, default=0.5)
    parser.add_argument(
        "--categories",
        nargs="+",
        default=None,
        help="Optional explicit label order. Defaults to first-seen order from the annotation file.",
    )
    parser.add_argument(
        "--skip-existing-videos",
        dest="skip_existing_videos",
        action="store_true",
        help=(
            "Skip a source video when an existing per-video manifest and outputs "
            "already match the current annotations/config. Default: enabled."
        ),
    )
    parser.add_argument(
        "--no-skip-existing-videos",
        dest="skip_existing_videos",
        action="store_false",
        help="Force rebuilding every referenced source video even if outputs already exist.",
    )
    parser.set_defaults(skip_existing_videos=True)
    return parser.parse_args()


def _read_xlsx_with_stdlib(path: Path, sheet_name: Optional[str]) -> List[List[object]]:
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}

    with zipfile.ZipFile(path) as zf:
        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall("a:si", ns):
                shared_strings.append("".join((t.text or "") for t in si.iterfind(".//a:t", ns)))

        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        sheets = []
        for sheet in workbook.find("a:sheets", ns):
            sheets.append(
                (
                    sheet.attrib.get("name"),
                    sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"),
                )
            )
        if not sheets:
            return []

        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rel_map = {
            rel.attrib["Id"]: rel.attrib["Target"] for rel in rels.findall("r:Relationship", rel_ns)
        }

        if sheet_name is None:
            target_rid = sheets[0][1]
        else:
            matched = [rid for name, rid in sheets if name == sheet_name]
            if not matched:
                raise KeyError(f"Sheet {sheet_name!r} not found in {path}")
            target_rid = matched[0]

        target = rel_map[target_rid]
        if not target.startswith("xl/"):
            target = f"xl/{target}"

        sheet_root = ET.fromstring(zf.read(target))
        rows = []
        for row in sheet_root.findall(".//a:row", ns):
            row_values: List[object] = []
            for cell in row.findall("a:c", ns):
                cell_type = cell.attrib.get("t")
                value_node = cell.find("a:v", ns)
                value = None if value_node is None else value_node.text
                if cell_type == "s" and value is not None:
                    value = shared_strings[int(value)]
                elif value is not None and cell_type != "str":
                    try:
                        numeric = float(value)
                        value = int(numeric) if numeric.is_integer() else numeric
                    except ValueError:
                        pass
                row_values.append(value)
            rows.append(row_values)
    return rows


def load_tabular_rows(path: Path, sheet_name: Optional[str]) -> List[Dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    try:
        import pandas as pd  # type: ignore
    except ModuleNotFoundError:
        pd = None

    if pd is not None:
        df = pd.read_excel(path, sheet_name=sheet_name)
        return df.where(df.notna(), None).to_dict(orient="records")

    try:
        from openpyxl import load_workbook  # type: ignore
    except ModuleNotFoundError:
        load_workbook = None

    if load_workbook is not None:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(item).strip() if item is not None else "" for item in rows[0]]
        output = []
        for values in rows[1:]:
            item = {}
            for key, value in zip(headers, values):
                item[key] = value
            output.append(item)
        return output

    rows = _read_xlsx_with_stdlib(path, sheet_name)
    if not rows:
        return []
    headers = [str(item).strip() if item is not None else "" for item in rows[0]]
    output = []
    for values in rows[1:]:
        item = {}
        for key, value in zip(headers, values):
            item[key] = value
        output.append(item)
    return output


def require_columns(rows: Sequence[Dict[str, object]], columns: Sequence[str]) -> None:
    if not rows:
        raise ValueError("Annotation file is empty.")
    header = set(rows[0].keys())
    missing = [col for col in columns if col and col not in header]
    if missing:
        raise KeyError(f"Missing required columns in annotation file: {missing}")


def normalize_video_key(value: object) -> str:
    if value is None:
        raise ValueError("Found empty video name in annotation file.")
    raw = str(value).strip().replace("\\", "/")
    if not raw:
        raise ValueError("Found blank video name in annotation file.")
    return raw


def to_int(value: object, field_name: str) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Missing value for {field_name}.")
    if isinstance(value, bool):
        raise ValueError(f"Invalid boolean value for {field_name}: {value}")
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Cannot parse {field_name}={value!r} as integer.") from exc


def sanitize_name(value: str) -> str:
    chars = []
    for char in value:
        if char.isalnum() or char in ("-", "_"):
            chars.append(char)
        else:
            chars.append("_")
    name = "".join(chars).strip("_")
    return name or "sample"


def build_video_index(video_root: Path) -> Dict[str, Path]:
    index: Dict[str, Path] = {}
    for video_path in sorted(video_root.rglob("*.mp4")):
        rel = video_path.relative_to(video_root)
        candidates = {
            rel.as_posix(),
            video_path.name,
            rel.stem,
            video_path.stem,
            rel.with_suffix("").as_posix(),
        }
        for key in candidates:
            index.setdefault(key, video_path)
    if not index:
        raise FileNotFoundError(f"No .mp4 files found under {video_root}")
    return index


def resolve_video_path(video_key: str, index: Dict[str, Path]) -> Path:
    video_key_no_ext = video_key[:-4] if video_key.lower().endswith(".mp4") else video_key
    candidates = [
        video_key,
        Path(video_key).name,
        Path(video_key).stem,
        video_key_no_ext,
        video_key.replace("\\", "/"),
    ]
    for key in candidates:
        if key in index:
            return index[key]
    raise FileNotFoundError(f"Cannot resolve video {video_key!r} under the provided --video-root.")


def build_segment_stem(label: str, start_frame: int, end_frame: int, duplicate_idx: int) -> str:
    base = f"{sanitize_name(label)}_f{start_frame:05d}_{end_frame:05d}"
    if duplicate_idx <= 1:
        return base
    return f"{base}_dup{duplicate_idx:02d}"


def assign_stable_segment_names(segments: List[Dict[str, object]]) -> None:
    duplicate_counts: Dict[Tuple[str, str, int, int], int] = {}
    for segment in segments:
        video_rel_stem = str(segment["video_rel_stem"])
        label = str(segment["label"])
        start_frame = int(segment["start_frame"])
        end_frame = int(segment["end_frame"])
        key = (video_rel_stem, label, start_frame, end_frame)
        duplicate_idx = duplicate_counts.get(key, 0) + 1
        duplicate_counts[key] = duplicate_idx
        segment_stem = build_segment_stem(label, start_frame, end_frame, duplicate_idx)
        segment["segment_stem"] = segment_stem
        segment["sample_name"] = f"{video_rel_stem}/{segment_stem}"
        segment["annotation_split"] = segment.get("split")


def group_segments_by_source_video(
    segments: Sequence[Dict[str, object]]
) -> "OrderedDict[str, List[Dict[str, object]]]":
    grouped: "OrderedDict[str, List[Dict[str, object]]]" = OrderedDict()
    for segment in segments:
        grouped.setdefault(str(segment["video_rel_stem"]), []).append(segment)
    return grouped


def build_prepare_config_signature(args: argparse.Namespace, categories: Sequence[str]) -> dict:
    signature = {
        "crop_first": bool(args.crop_first),
        "frame_base": int(args.frame_base),
        "pre_pad": int(args.pre_pad),
        "post_pad": int(args.post_pad),
        "min_frames": int(args.min_frames),
        "min_valid_frames": int(args.min_valid_frames),
        "min_valid_ratio": float(args.min_valid_ratio),
        "handedness": str(args.handedness),
        "track_id": args.track_id,
        "use_smooth": bool(args.use_smooth),
        "joint_score_source": str(args.joint_score_source),
        "model_path": str(args.model_path),
        "num_hands": int(args.num_hands),
        "padding_ratio": float(args.padding_ratio),
        "min_hand_detection_confidence": float(args.min_hand_detection_confidence),
        "min_hand_presence_confidence": float(args.min_hand_presence_confidence),
        "min_tracking_confidence": float(args.min_tracking_confidence),
        "categories": list(categories),
    }
    optional_keys = [
        "low_quality_valid_ratio",
        "other_track_coverage_threshold",
        "max_consecutive_missing_threshold",
        "max_interpolation_gap",
        "skip_quality_filter",
        "prepare_mode",
    ]
    for key in optional_keys:
        if hasattr(args, key):
            signature[key] = getattr(args, key)
    return signature


def build_source_file_signature(source_video_path: Path) -> dict:
    stat = source_video_path.stat()
    return {
        "path": str(source_video_path),
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


def build_video_signature(
    args: argparse.Namespace,
    source_video_path: Path,
    video_rel_stem: str,
    segments: Sequence[Dict[str, object]],
    categories: Sequence[str],
) -> dict:
    return {
        "source_video": build_source_file_signature(source_video_path),
        "source_video_rel_stem": video_rel_stem,
        "config": build_prepare_config_signature(args, categories),
        "segments": [
            {
                "row_idx": int(segment["row_idx"]),
                "label": str(segment["label"]),
                "start_frame": int(segment["start_frame"]),
                "end_frame": int(segment["end_frame"]),
                "segment_stem": str(segment["segment_stem"]),
            }
            for segment in segments
        ],
    }


def get_video_output_dir(output_root: Path, video_rel_stem: str) -> Path:
    return output_root / video_rel_stem


def get_video_clips_dir(output_root: Path, video_rel_stem: str) -> Path:
    return get_video_output_dir(output_root, video_rel_stem) / "clips"


def get_video_manifest_path(output_root: Path, video_rel_stem: str) -> Path:
    return get_video_output_dir(output_root, video_rel_stem) / "video_manifest.json"


def get_video_crop_dir(crop_root: Path, video_rel_stem: str) -> Path:
    return crop_root / video_rel_stem


def get_video_kpts_dir(kpts_root: Path, video_rel_stem: str) -> Path:
    return kpts_root / video_rel_stem


def segment_clip_json_path(output_root: Path, video_rel_stem: str, segment_stem: str) -> Path:
    return get_video_clips_dir(output_root, video_rel_stem) / f"{segment_stem}.json"


def segment_crop_video_path(crop_root: Path, video_rel_stem: str, segment_stem: str) -> Path:
    return get_video_crop_dir(crop_root, video_rel_stem) / f"{segment_stem}.mp4"


def segment_kpts_annotations_path(kpts_root: Path, video_rel_stem: str, segment_stem: str) -> Path:
    return get_video_kpts_dir(kpts_root, video_rel_stem) / segment_stem / "annotations.json"


def load_json_if_exists(path: Path) -> Optional[dict]:
    if not path.exists():
        return None
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def remove_path_if_exists(path: Path) -> None:
    if not path.exists():
        return
    if path.is_dir():
        shutil.rmtree(path)
    else:
        path.unlink()


def video_outputs_complete(
    args: argparse.Namespace,
    output_root: Path,
    kpts_root: Path,
    crop_root: Path,
    video_rel_stem: str,
    segments: Sequence[Dict[str, object]],
) -> bool:
    manifest_path = get_video_manifest_path(output_root, video_rel_stem)
    if not manifest_path.exists():
        return False

    for segment in segments:
        segment_stem = str(segment["segment_stem"])
        clip_json_path = segment_clip_json_path(output_root, video_rel_stem, segment_stem)
        if not clip_json_path.exists():
            return False
        if args.crop_first:
            if not segment_kpts_annotations_path(kpts_root, video_rel_stem, segment_stem).exists():
                return False
        else:
            if not (get_video_kpts_dir(kpts_root, video_rel_stem) / "annotations.json").exists():
                return False
    return True


def load_compact_4col_rows(path: Path, sheet_name: Optional[str]) -> List[Dict[str, object]]:
    suffix = path.suffix.lower()
    raw_rows: List[List[object]]

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            raw_rows = [row for row in csv.reader(f)]
    else:
        try:
            import pandas as pd  # type: ignore
        except ModuleNotFoundError:
            pd = None

        if pd is not None:
            df = pd.read_excel(path, sheet_name=sheet_name, header=None)
            raw_rows = df.where(df.notna(), None).values.tolist()
        else:
            try:
                from openpyxl import load_workbook  # type: ignore
            except ModuleNotFoundError:
                load_workbook = None

            if load_workbook is not None:
                wb = load_workbook(path, read_only=True, data_only=True)
                ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
                raw_rows = [list(row) for row in ws.iter_rows(values_only=True)]
            else:
                raw_rows = _read_xlsx_with_stdlib(path, sheet_name)

    output = []
    current_video = None
    for row_idx, row in enumerate(raw_rows, start=1):
        if not row:
            continue
        padded = list(row) + [None] * (4 - len(row))
        video_value, start_value, end_value, label_value = padded[:4]
        if video_value not in (None, ""):
            current_video = str(video_value).strip()
        if current_video is None:
            raise ValueError(f"Compact annotation row {row_idx} is missing a video name before segment data.")
        if start_value in (None, "") and end_value in (None, "") and label_value in (None, ""):
            continue
        output.append(
            {
                "video_name": current_video,
                "start_frame": start_value,
                "end_frame": end_value,
                "label": label_value,
            }
        )
    return output


def choose_track(
    frames: Sequence[dict],
    preferred_handedness: str,
    explicit_track_id: Optional[int],
) -> int:
    stats: Dict[int, Dict[str, object]] = {}
    for frame in frames:
        for hand in frame.get("hands", []):
            track_id = int(hand["track_id"])
            handedness = str(hand.get("handedness", "Unknown"))
            score = float(hand.get("score", 0.0))
            item = stats.setdefault(track_id, {"count": 0, "score_sum": 0.0, "handedness": {}})
            item["count"] = int(item["count"]) + 1
            item["score_sum"] = float(item["score_sum"]) + score
            handedness_counter = item["handedness"]
            handedness_counter[handedness] = handedness_counter.get(handedness, 0) + 1

    if explicit_track_id is not None:
        if explicit_track_id not in stats:
            raise ValueError(f"track_id={explicit_track_id} not found in annotations.")
        return explicit_track_id

    if not stats:
        raise ValueError("No hand tracks found in annotations.")

    def sort_key(track_id: int) -> Tuple[int, int, float]:
        item = stats[track_id]
        handedness_counter = item["handedness"]
        dominant = (
            max(handedness_counter.items(), key=lambda kv: kv[1])[0]
            if handedness_counter
            else "Unknown"
        )
        handedness_match = int(preferred_handedness != "any" and dominant == preferred_handedness)
        return handedness_match, int(item["count"]), float(item["score_sum"])

    return max(stats.keys(), key=sort_key)


def select_segment_track(
    annotations: dict,
    start_frame: int,
    end_frame: int,
    preferred_handedness: str,
    explicit_track_id: Optional[int],
) -> int:
    segment_frames = []
    for frame in annotations.get("frames", []):
        frame_idx = int(frame.get("frame_idx", -1))
        if start_frame <= frame_idx <= end_frame:
            segment_frames.append(frame)
    return choose_track(segment_frames, preferred_handedness, explicit_track_id)


def extract_hand_points(hand: dict, use_smooth: bool) -> np.ndarray:
    coords = hand["kpts_xy_smooth"] if use_smooth and "kpts_xy_smooth" in hand else hand["kpts_xy"]
    arr = np.asarray(coords, dtype=np.float32)
    if arr.shape != (21, 2):
        raise ValueError(f"Expected hand keypoints shape (21, 2), got {arr.shape}")
    return arr


def make_joint_score(hand: dict, source: str) -> float:
    if source == "ones":
        return 1.0
    return float(hand.get("score", 0.0))


def build_clip_json(
    annotations: dict,
    track_id: int,
    start_frame: int,
    end_frame: int,
    sample_name: str,
    category_id: int,
    use_smooth: bool,
    joint_score_source: str,
) -> dict:
    resolution = annotations.get("resolution")
    if not resolution or len(resolution) != 2:
        raise ValueError("annotations.json missing valid resolution")
    width, height = int(resolution[0]), int(resolution[1])
    frame_map = {int(frame["frame_idx"]): frame for frame in annotations.get("frames", [])}
    clip_frames = end_frame - start_frame + 1
    clip_ann = []

    for offset, frame_idx in enumerate(range(start_frame, end_frame + 1)):
        frame = frame_map.get(frame_idx, {})
        points = np.zeros((21, 2), dtype=np.float32)
        score = 0.0
        for hand in frame.get("hands", []):
            if int(hand["track_id"]) != track_id:
                continue
            points = extract_hand_points(hand, use_smooth)
            score = make_joint_score(hand, joint_score_source)
            break
        clip_ann.append(
            {
                "frame_index": offset,
                "source_frame_index": frame_idx,
                "keypoints": [[float(x), float(y), float(score)] for x, y in points],
            }
        )

    return {
        "info": {
            "video_name": sample_name,
            "source_video": annotations.get("video", ""),
            "resolution": [width, height],
            "num_frame": clip_frames,
            "num_keypoints": 21,
            "source_start_frame": start_frame,
            "source_end_frame": end_frame,
            "track_id": track_id,
        },
        "category_id": category_id,
        "annotations": clip_ann,
    }


def compute_clip_valid_stats(
    annotations: dict,
    track_id: int,
    start_frame: int,
    end_frame: int,
) -> Dict[str, float]:
    frame_map = {int(frame["frame_idx"]): frame for frame in annotations.get("frames", [])}
    clip_frames = end_frame - start_frame + 1
    valid_frames = 0
    score_sum = 0.0

    for frame_idx in range(start_frame, end_frame + 1):
        frame = frame_map.get(frame_idx, {})
        for hand in frame.get("hands", []):
            if int(hand["track_id"]) != track_id:
                continue
            valid_frames += 1
            score_sum += float(hand.get("score", 0.0))
            break

    valid_ratio = valid_frames / clip_frames if clip_frames > 0 else 0.0
    mean_hand_score = score_sum / valid_frames if valid_frames > 0 else 0.0
    return {
        "clip_frames": float(clip_frames),
        "valid_frames": float(valid_frames),
        "valid_ratio": float(valid_ratio),
        "mean_hand_score": float(mean_hand_score),
    }


def load_extractor_module():
    spec = importlib.util.spec_from_file_location("extract_hand_roi_with_kpts", EXTRACTOR_PATH)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot load extractor module from {EXTRACTOR_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def maybe_run_extraction(args: argparse.Namespace, kpts_root: Path, src_root: Optional[Path] = None) -> None:
    if args.skip_extract:
        log_status(f"[extract] Skip extraction and reuse existing annotations under {kpts_root}")
        return
    extract_src_root = src_root or args.video_root
    video_count = sum(1 for _ in extract_src_root.rglob("*.mp4"))
    log_status(
        f"[extract] Start MediaPipe extraction from {extract_src_root} "
        f"to {kpts_root} ({video_count} videos)"
    )
    module = load_extractor_module()
    video_keys = [
        str(path.relative_to(extract_src_root).as_posix())
        for path in sorted(extract_src_root.rglob("*.mp4"))
    ]
    extract_videos_with_list(
        module=module,
        src_dir=extract_src_root,
        dst_dir=kpts_root,
        video_keys=video_keys,
        args=args,
        save_vis=False,
    )
    log_status(f"[extract] Finished MediaPipe extraction into {kpts_root}")


def build_detector_kwargs(args: argparse.Namespace) -> dict:
    return dict(
        model_path=args.model_path,
        num_hands=args.num_hands,
        min_hand_detection_confidence=args.min_hand_detection_confidence,
        min_hand_presence_confidence=args.min_hand_presence_confidence,
        min_tracking_confidence=args.min_tracking_confidence,
    )


def extract_videos_with_list(
    *,
    module,
    src_dir: Path,
    dst_dir: Path,
    video_keys: Sequence[str],
    args: argparse.Namespace,
    save_vis: bool = False,
    max_workers: Optional[int] = None,
) -> None:
    if not video_keys:
        return

    fd, tmp_path = tempfile.mkstemp(prefix="hand21_video_list_", suffix=".txt")
    list_path = Path(tmp_path)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            for key in video_keys:
                f.write(f"{key}\n")
        module.extract_hand_roi_from_videos(
            src_dir=str(src_dir),
            dst_dir=str(dst_dir),
            video_list_file=str(list_path),
            model_path=args.model_path,
            padding_ratio=args.padding_ratio,
            num_hands=args.num_hands,
            min_hand_detection_confidence=args.min_hand_detection_confidence,
            min_hand_presence_confidence=args.min_hand_presence_confidence,
            min_tracking_confidence=args.min_tracking_confidence,
            save_vis=save_vis,
            max_workers=max_workers,
        )
    finally:
        list_path.unlink(missing_ok=True)


def extract_single_source_video(
    module,
    args: argparse.Namespace,
    source_video_path: Path,
    video_kpts_dir: Path,
    log_prefix: str,
) -> None:
    if args.skip_extract and (video_kpts_dir / "annotations.json").exists():
        log_status(f"{log_prefix} Reuse existing source-video MediaPipe annotations: {video_kpts_dir}")
        return

    detector = module.MediaPipeHandDetector(**build_detector_kwargs(args))
    tracker = module.HandTracker()
    log_status(f"{log_prefix} MediaPipe source video -> {video_kpts_dir}")
    try:
        module.process_single_video(
            video_path=source_video_path,
            output_dir=video_kpts_dir,
            detector=detector,
            padding_ratio=args.padding_ratio,
            tracker=tracker,
            vis_path=None,
        )
    finally:
        detector.close()


def extract_video_clips(
    module,
    args: argparse.Namespace,
    video_crop_dir: Path,
    video_kpts_dir: Path,
    log_prefix: str,
) -> None:
    if args.skip_extract and all((video_kpts_dir / clip.stem / "annotations.json").exists() for clip in video_crop_dir.glob("*.mp4")):
        log_status(f"{log_prefix} Reuse existing clip-level MediaPipe annotations: {video_kpts_dir}")
        return

    log_status(f"{log_prefix} MediaPipe cropped clips -> {video_kpts_dir}")
    video_keys = [clip.name for clip in sorted(video_crop_dir.glob("*.mp4"))]
    extract_videos_with_list(
        module=module,
        src_dir=video_crop_dir,
        dst_dir=video_kpts_dir,
        video_keys=video_keys,
        args=args,
        save_vis=False,
        max_workers=1,
    )


def crop_source_video_segments(
    segments: Sequence[Dict[str, object]],
    video_rel_stem: str,
    crop_root: Path,
    log_prefix: str,
) -> Dict[str, int]:
    try:
        import cv2  # type: ignore
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(
            "OpenCV (cv2) is required for --crop-first. Please install opencv-python in this environment."
        ) from exc

    video_crop_dir = get_video_crop_dir(crop_root, video_rel_stem)
    video_crop_dir.mkdir(parents=True, exist_ok=True)
    source_video = Path(str(segments[0]["video_path"]))
    cap = cv2.VideoCapture(str(source_video))
    if not cap.isOpened():
        raise RuntimeError(f"Cannot open source video for cropping: {source_video}")
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    cap.release()

    written = 0
    skipped = 0
    total_segments = len(segments)
    log_status(f"{log_prefix} Crop {total_segments} segments into {video_crop_dir}")

    for seg_idx, segment in enumerate(segments, start=1):
        start_frame = max(0, int(segment["start_frame"]))
        end_frame = min(total_frames - 1, int(segment["end_frame"]))
        clip_frames = end_frame - start_frame + 1
        if clip_frames <= 0:
            segment["skip_reason"] = "out_of_range_after_clamp"
            skipped += 1
            if should_report_progress(seg_idx, total_segments):
                log_status(
                    f"{log_prefix} Crop progress {format_progress(seg_idx, total_segments)} | "
                    f"written={written} skipped={skipped}"
                )
            continue

        clip_out_path = segment_crop_video_path(crop_root, video_rel_stem, str(segment["segment_stem"]))
        clip_out_path.parent.mkdir(parents=True, exist_ok=True)

        cap = cv2.VideoCapture(str(source_video))
        if not cap.isOpened():
            raise RuntimeError(f"Cannot open source video for cropping: {source_video}")

        fps = float(cap.get(cv2.CAP_PROP_FPS))
        if not fps or fps <= 0:
            fps = 30.0
        width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        if width <= 0 or height <= 0:
            cap.release()
            raise RuntimeError(f"Invalid source resolution for cropping: {source_video}")

        writer = cv2.VideoWriter(
            str(clip_out_path),
            cv2.VideoWriter_fourcc(*"mp4v"),
            fps,
            (width, height),
        )
        if not writer.isOpened():
            cap.release()
            raise RuntimeError(f"Cannot open output clip for writing: {clip_out_path}")

        cap.set(cv2.CAP_PROP_POS_FRAMES, start_frame)
        copied = 0
        frame_idx = start_frame
        while frame_idx <= end_frame:
            ok, frame = cap.read()
            if not ok:
                break
            writer.write(frame)
            copied += 1
            frame_idx += 1

        cap.release()
        writer.release()

        if copied <= 0:
            if clip_out_path.exists():
                clip_out_path.unlink()
            segment["skip_reason"] = "crop_decode_failed"
            skipped += 1
            if should_report_progress(seg_idx, total_segments):
                log_status(
                    f"{log_prefix} Crop progress {format_progress(seg_idx, total_segments)} | "
                    f"written={written} skipped={skipped}"
                )
            continue

        segment["crop_num_frames"] = copied
        segment["crop_source_start_frame"] = start_frame
        segment["crop_source_end_frame"] = end_frame
        written += 1
        if should_report_progress(seg_idx, total_segments):
            log_status(
                f"{log_prefix} Crop progress {format_progress(seg_idx, total_segments)} | "
                f"written={written} skipped={skipped}"
            )

    log_status(f"{log_prefix} Crop finished | written={written} skipped={skipped}")
    return {
        "num_segments_cropped": written,
        "num_segments_crop_skipped": skipped,
    }


def crop_segment_videos(
    segments: List[Dict[str, object]],
    crop_root: Path,
) -> Dict[str, int]:
    try:
        import cv2  # type: ignore
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(
            "OpenCV (cv2) is required for --crop-first. Please install opencv-python in this environment."
        ) from exc

    crop_root.mkdir(parents=True, exist_ok=True)
    video_frame_count_cache: Dict[Path, int] = {}
    written = 0
    skipped = 0
    total_segments = len(segments)
    log_status(f"[crop] Start cropping {total_segments} annotated segments into {crop_root}")

    for seg_idx, segment in enumerate(segments):
        current = seg_idx + 1
        source_video = Path(str(segment["video_path"]))
        if source_video not in video_frame_count_cache:
            cap = cv2.VideoCapture(str(source_video))
            if not cap.isOpened():
                raise RuntimeError(f"Cannot open source video for cropping: {source_video}")
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            cap.release()
            video_frame_count_cache[source_video] = total_frames

        total_frames = video_frame_count_cache[source_video]
        start_frame = max(0, int(segment["start_frame"]))
        end_frame = min(total_frames - 1, int(segment["end_frame"]))
        clip_frames = end_frame - start_frame + 1
        if clip_frames <= 0:
            segment["skip_reason"] = "out_of_range_after_clamp"
            skipped += 1
            if should_report_progress(current, total_segments):
                log_status(
                    f"[crop] Progress {format_progress(current, total_segments)} | "
                    f"written={written} skipped={skipped}"
                )
            continue

        rel_stem = Path(str(segment["video_rel_stem"]))
        label = sanitize_name(str(segment["label"]))
        clip_stem = f"{rel_stem.stem}_seg{seg_idx:04d}_{label}_f{start_frame:05d}_{end_frame:05d}"
        clip_rel_stem = (rel_stem.parent / clip_stem).as_posix()
        clip_out_path = crop_root / f"{clip_rel_stem}.mp4"
        clip_out_path.parent.mkdir(parents=True, exist_ok=True)

        cap = cv2.VideoCapture(str(source_video))
        if not cap.isOpened():
            raise RuntimeError(f"Cannot open source video for cropping: {source_video}")

        fps = float(cap.get(cv2.CAP_PROP_FPS))
        if not fps or fps <= 0:
            fps = 30.0
        width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        if width <= 0 or height <= 0:
            cap.release()
            raise RuntimeError(f"Invalid source resolution for cropping: {source_video}")

        writer = cv2.VideoWriter(
            str(clip_out_path),
            cv2.VideoWriter_fourcc(*"mp4v"),
            fps,
            (width, height),
        )
        if not writer.isOpened():
            cap.release()
            raise RuntimeError(f"Cannot open output clip for writing: {clip_out_path}")

        cap.set(cv2.CAP_PROP_POS_FRAMES, start_frame)
        copied = 0
        frame_idx = start_frame
        while frame_idx <= end_frame:
            ok, frame = cap.read()
            if not ok:
                break
            writer.write(frame)
            copied += 1
            frame_idx += 1

        cap.release()
        writer.release()

        if copied <= 0:
            if clip_out_path.exists():
                clip_out_path.unlink()
            segment["skip_reason"] = "crop_decode_failed"
            skipped += 1
            if should_report_progress(current, total_segments):
                log_status(
                    f"[crop] Progress {format_progress(current, total_segments)} | "
                    f"written={written} skipped={skipped}"
                )
            continue

        segment["crop_rel_stem"] = clip_rel_stem
        segment["crop_num_frames"] = copied
        segment["crop_source_start_frame"] = start_frame
        segment["crop_source_end_frame"] = end_frame
        written += 1
        if should_report_progress(current, total_segments):
            log_status(
                f"[crop] Progress {format_progress(current, total_segments)} | "
                f"written={written} skipped={skipped}"
            )

    log_status(f"[crop] Finished cropping: written={written}, skipped={skipped}, output={crop_root}")
    return {
        "num_segments_cropped": written,
        "num_segments_crop_skipped": skipped,
    }


def normalize_split_name(value: object) -> str:
    if value is None:
        raise ValueError("Split value is empty.")
    name = str(value).strip().lower()
    alias = {
        "training": "train",
        "train": "train",
        "val": "val",
        "valid": "val",
        "validation": "val",
        "test": "test",
        "testing": "test",
    }
    if name not in alias:
        raise ValueError(f"Unsupported split value: {value!r}")
    return alias[name]


def grouped_random_split(
    videos: Sequence[str],
    val_ratio: float,
    test_ratio: float,
    seed: int,
) -> Dict[str, str]:
    if val_ratio < 0 or test_ratio < 0 or val_ratio + test_ratio >= 1:
        raise ValueError("val_ratio and test_ratio must be >= 0 and sum to < 1.")
    shuffled = list(videos)
    rng = random.Random(seed)
    rng.shuffle(shuffled)
    total = len(shuffled)
    val_count = int(round(total * val_ratio))
    test_count = int(round(total * test_ratio))
    if total > 1 and val_ratio > 0 and val_count == 0:
        val_count = 1
    if total - val_count > 1 and test_ratio > 0 and test_count == 0:
        test_count = 1
    if val_count + test_count >= total and total > 1:
        test_count = max(0, min(test_count, total - 1))
        val_count = max(1 if val_ratio > 0 else 0, min(val_count, total - test_count - 1))

    split_map: Dict[str, str] = {}
    for idx, video_name in enumerate(shuffled):
        if idx < val_count:
            split_map[video_name] = "val"
        elif idx < val_count + test_count:
            split_map[video_name] = "test"
        else:
            split_map[video_name] = "train"
    return split_map


def collect_segments(
    args: argparse.Namespace,
    rows: Sequence[Dict[str, object]],
    video_index: Dict[str, Path],
) -> Tuple[List[Dict[str, object]], List[str], List[Dict[str, object]], List[Dict[str, object]]]:
    if args.annotation_format == "table":
        required = [args.video_col, args.label_col, args.start_col, args.end_col]
        if args.split_col:
            required.append(args.split_col)
        require_columns(rows, required)

    categories_ordered = OrderedDict()
    segments: List[Dict[str, object]] = []
    missing_video_rows: List[Dict[str, object]] = []
    invalid_frame_rows: List[Dict[str, object]] = []
    total_rows = len(rows)
    log_status(f"[collect] Validate and resolve {total_rows} annotation rows")

    for processed_idx, row in enumerate(rows, start=1):
        row_idx = processed_idx + 1
        video_key = normalize_video_key(row[args.video_col])
        label = str(row[args.label_col]).strip()
        if not label:
            raise ValueError(f"Row {row_idx}: empty label.")
        start_frame = to_int(row[args.start_col], args.start_col)
        end_frame = to_int(row[args.end_col], args.end_col)
        if args.frame_base == 1:
            start_frame -= 1
            end_frame -= 1
        start_frame -= args.pre_pad
        end_frame += args.post_pad
        if end_frame < start_frame:
            invalid_frame_rows.append(
                {
                    "row_idx": row_idx,
                    "video_key": video_key,
                    "label": label,
                    "start_frame": start_frame,
                    "end_frame": end_frame,
                    "reason": "end frame earlier than start frame",
                }
            )
            if should_report_progress(processed_idx, total_rows):
                log_status(
                    f"[collect] Progress {format_progress(processed_idx, total_rows)} | "
                    f"valid={len(segments)} missing_video={len(missing_video_rows)} "
                    f"invalid_frame={len(invalid_frame_rows)}"
                )
            continue

        try:
            source_video_path = resolve_video_path(video_key, video_index)
        except FileNotFoundError:
            missing_video_rows.append(
                {
                    "row_idx": row_idx,
                    "video_key": video_key,
                    "label": label,
                    "start_frame": start_frame,
                    "end_frame": end_frame,
                }
            )
            if should_report_progress(processed_idx, total_rows):
                log_status(
                    f"[collect] Progress {format_progress(processed_idx, total_rows)} | "
                    f"valid={len(segments)} missing_video={len(missing_video_rows)} "
                    f"invalid_frame={len(invalid_frame_rows)}"
                )
            continue
        rel_video = source_video_path.relative_to(args.video_root).with_suffix("").as_posix()
        segment = {
            "row_idx": row_idx,
            "video_key": video_key,
            "video_rel_stem": rel_video,
            "video_path": source_video_path,
            "label": label,
            "start_frame": start_frame,
            "end_frame": end_frame,
            "split": row[args.split_col] if args.split_col else None,
        }
        segments.append(segment)
        categories_ordered.setdefault(label, None)
        if should_report_progress(processed_idx, total_rows):
            log_status(
                f"[collect] Progress {format_progress(processed_idx, total_rows)} | "
                f"valid={len(segments)} missing_video={len(missing_video_rows)} "
                f"invalid_frame={len(invalid_frame_rows)}"
            )

    categories = args.categories or list(categories_ordered.keys())
    unknown_labels = sorted({seg["label"] for seg in segments} - set(categories))
    if unknown_labels:
        raise ValueError(f"Labels found in annotations but missing from --categories: {unknown_labels}")
    log_status(
        f"[collect] Finished: valid_segments={len(segments)}, categories={len(categories)}, "
        f"missing_video_rows={len(missing_video_rows)}, invalid_frame_rows={len(invalid_frame_rows)}"
    )
    return segments, categories, missing_video_rows, invalid_frame_rows


def ensure_split_assignment(
    segments: List[Dict[str, object]],
    val_ratio: float,
    test_ratio: float,
    seed: int,
) -> None:
    if all(seg["split"] is not None for seg in segments):
        return
    videos = sorted({str(seg["video_rel_stem"]) for seg in segments})
    split_map = grouped_random_split(videos, val_ratio, test_ratio, seed)
    for seg in segments:
        seg["split"] = split_map[str(seg["video_rel_stem"])]


def main() -> None:
    args = parse_args()
    kpts_root = args.kpts_root or args.output_root.parent / f"{args.output_root.name}_video_kpts"
    crop_root = args.crop_root or args.output_root.parent / f"{args.output_root.name}_segments_videos"
    log_status(
        f"[start] video_root={args.video_root} | annotation_file={args.annotation_file} | "
        f"output_root={args.output_root} | crop_first={args.crop_first} | "
        f"skip_existing_videos={args.skip_existing_videos}"
    )
    log_status("[mode] This preparation script no longer performs train/val/test splitting.")
    if args.split_col:
        log_status(f"[mode] split_col={args.split_col} will only be preserved in metadata.")
    if args.val_ratio != 0.2 or args.test_ratio != 0.0 or args.seed != 42:
        log_status("[mode] val_ratio/test_ratio/seed are ignored during preparation.")

    video_index = build_video_index(args.video_root)
    log_status(f"[scan] Indexed {len(video_index)} video lookup keys under {args.video_root}")

    if args.annotation_format == "compact_4col":
        rows = load_compact_4col_rows(args.annotation_file, args.sheet_name)
    else:
        rows = load_tabular_rows(args.annotation_file, args.sheet_name)
        if args.annotation_format == "auto":
            required = {args.video_col, args.label_col, args.start_col, args.end_col}
            header = set(rows[0].keys()) if rows else set()
            if not required.issubset(header):
                rows = load_compact_4col_rows(args.annotation_file, args.sheet_name)
                args.annotation_format = "compact_4col"
            else:
                args.annotation_format = "table"
    log_status(f"[load] Loaded {len(rows)} annotation rows with format={args.annotation_format}")

    segments, categories, missing_video_rows, invalid_frame_rows = collect_segments(args, rows, video_index)
    if not segments:
        missing_keys = sorted({str(item["video_key"]) for item in missing_video_rows})
        invalid_keys = sorted({str(item["video_key"]) for item in invalid_frame_rows})
        preview_keys = missing_keys + [key for key in invalid_keys if key not in missing_keys]
        preview = ", ".join(preview_keys[:10])
        if len(preview_keys) > 10:
            preview += ", ..."
        raise FileNotFoundError(
            "No valid annotation rows remained after video/frame validation. "
            f"Affected video keys: {preview or 'none'}"
        )
    assign_stable_segment_names(segments)
    segments_by_video = group_segments_by_source_video(segments)
    log_status(
        f"[group] Prepared {len(segments)} valid segments across {len(segments_by_video)} source videos"
    )

    label_to_id = {label: idx for idx, label in enumerate(categories)}
    args.output_root.mkdir(parents=True, exist_ok=True)
    kpts_root.mkdir(parents=True, exist_ok=True)
    if args.crop_first:
        crop_root.mkdir(parents=True, exist_ok=True)

    extractor_module = None if args.skip_extract else load_extractor_module()
    clip_meta: "OrderedDict[str, Dict[str, object]]" = OrderedDict()
    written = 0
    skipped_short = 0
    skipped_low_valid = 0
    processed_videos = 0
    skipped_videos = 0
    rebuilt_videos = 0
    crop_written = 0
    crop_skipped = 0
    total_videos = len(segments_by_video)

    for video_idx, (video_rel_stem, video_segments) in enumerate(segments_by_video.items(), start=1):
        source_video_path = Path(str(video_segments[0]["video_path"]))
        log_prefix = f"[video {video_idx}/{total_videos}] {video_rel_stem} |"
        current_signature = build_video_signature(
            args=args,
            source_video_path=source_video_path,
            video_rel_stem=video_rel_stem,
            segments=video_segments,
            categories=categories,
        )
        manifest_path = get_video_manifest_path(args.output_root, video_rel_stem)
        existing_manifest = load_json_if_exists(manifest_path)

        if (
            args.skip_existing_videos
            and existing_manifest is not None
            and existing_manifest.get("signature") == current_signature
            and video_outputs_complete(
                args=args,
                output_root=args.output_root,
                kpts_root=kpts_root,
                crop_root=crop_root,
                video_rel_stem=video_rel_stem,
                segments=video_segments,
            )
        ):
            skipped_videos += 1
            log_status(f"{log_prefix} Skip unchanged video with {len(video_segments)} segments")
            for sample_name, meta in existing_manifest.get("clips", {}).items():
                clip_meta[sample_name] = meta
            continue

        video_output_dir = get_video_output_dir(args.output_root, video_rel_stem)
        video_clips_dir = get_video_clips_dir(args.output_root, video_rel_stem)
        video_kpts_dir = get_video_kpts_dir(kpts_root, video_rel_stem)
        video_crop_dir = get_video_crop_dir(crop_root, video_rel_stem)

        if existing_manifest is not None or video_output_dir.exists() or video_kpts_dir.exists() or (
            args.crop_first and video_crop_dir.exists()
        ):
            rebuilt_videos += 1
            log_status(f"{log_prefix} Rebuild outputs for this video because inputs or config changed")

        processed_videos += 1
        remove_path_if_exists(video_output_dir)
        if args.crop_first:
            remove_path_if_exists(video_crop_dir)
        if not args.skip_extract:
            remove_path_if_exists(video_kpts_dir)

        video_clips_dir.mkdir(parents=True, exist_ok=True)

        video_crop_stats = {
            "num_segments_cropped": 0,
            "num_segments_crop_skipped": 0,
        }
        if args.crop_first:
            video_crop_stats = crop_source_video_segments(
                segments=video_segments,
                video_rel_stem=video_rel_stem,
                crop_root=crop_root,
                log_prefix=log_prefix,
            )
            crop_written += int(video_crop_stats["num_segments_cropped"])
            crop_skipped += int(video_crop_stats["num_segments_crop_skipped"])
            if extractor_module is not None:
                extract_video_clips(
                    module=extractor_module,
                    args=args,
                    video_crop_dir=video_crop_dir,
                    video_kpts_dir=video_kpts_dir,
                    log_prefix=log_prefix,
                )
        elif extractor_module is not None:
            extract_single_source_video(
                module=extractor_module,
                args=args,
                source_video_path=source_video_path,
                video_kpts_dir=video_kpts_dir,
                log_prefix=log_prefix,
            )

        video_clip_meta: "OrderedDict[str, Dict[str, object]]" = OrderedDict()
        video_written = 0
        video_skipped_short = 0
        video_skipped_low_valid = 0
        total_video_segments = len(video_segments)
        log_status(f"{log_prefix} Build hand21 JSON clips from {total_video_segments} segments")

        source_video_annotations = None
        if not args.crop_first:
            source_anno_path = video_kpts_dir / "annotations.json"
            if not source_anno_path.exists():
                raise FileNotFoundError(f"{log_prefix} Missing MediaPipe annotations: {source_anno_path}")
            with source_anno_path.open("r", encoding="utf-8") as f:
                source_video_annotations = json.load(f)

        for seg_idx, segment in enumerate(video_segments, start=1):
            segment_stem = str(segment["segment_stem"])
            current = seg_idx

            if segment.get("skip_reason"):
                video_skipped_short += 1
                if should_report_progress(current, total_video_segments):
                    log_status(
                        f"{log_prefix} Clip progress {format_progress(current, total_video_segments)} | "
                        f"written={video_written} skipped_short={video_skipped_short} "
                        f"skipped_low_valid={video_skipped_low_valid}"
                    )
                continue

            if args.crop_first:
                anno_path = segment_kpts_annotations_path(kpts_root, video_rel_stem, segment_stem)
                if not anno_path.exists():
                    raise FileNotFoundError(f"{log_prefix} Missing MediaPipe annotations for clip: {anno_path}")
                with anno_path.open("r", encoding="utf-8") as f:
                    annotations = json.load(f)
                total_frames = int(annotations.get("total_frames", 0))
                start_frame = 0
                end_frame = total_frames - 1
                source_start_frame = int(segment.get("crop_source_start_frame", segment["start_frame"]))
                source_end_frame = int(segment.get("crop_source_end_frame", segment["end_frame"]))
                kpts_source = f"{video_rel_stem}/{segment_stem}"
            else:
                annotations = source_video_annotations
                total_frames = int(annotations.get("total_frames", 0))
                start_frame = max(0, int(segment["start_frame"]))
                end_frame = min(total_frames - 1, int(segment["end_frame"]))
                source_start_frame = start_frame
                source_end_frame = end_frame
                kpts_source = video_rel_stem

            clip_frames = end_frame - start_frame + 1
            if clip_frames < args.min_frames:
                video_skipped_short += 1
                if should_report_progress(current, total_video_segments):
                    log_status(
                        f"{log_prefix} Clip progress {format_progress(current, total_video_segments)} | "
                        f"written={video_written} skipped_short={video_skipped_short} "
                        f"skipped_low_valid={video_skipped_low_valid}"
                    )
                continue

            try:
                track_id = select_segment_track(
                    annotations=annotations,
                    start_frame=start_frame,
                    end_frame=end_frame,
                    preferred_handedness=args.handedness,
                    explicit_track_id=args.track_id,
                )
            except ValueError:
                video_skipped_low_valid += 1
                if should_report_progress(current, total_video_segments):
                    log_status(
                        f"{log_prefix} Clip progress {format_progress(current, total_video_segments)} | "
                        f"written={video_written} skipped_short={video_skipped_short} "
                        f"skipped_low_valid={video_skipped_low_valid}"
                    )
                continue

            valid_stats = compute_clip_valid_stats(
                annotations=annotations,
                track_id=track_id,
                start_frame=start_frame,
                end_frame=end_frame,
            )
            if (
                int(valid_stats["valid_frames"]) < args.min_valid_frames
                or float(valid_stats["valid_ratio"]) < args.min_valid_ratio
            ):
                video_skipped_low_valid += 1
                if should_report_progress(current, total_video_segments):
                    log_status(
                        f"{log_prefix} Clip progress {format_progress(current, total_video_segments)} | "
                        f"written={video_written} skipped_short={video_skipped_short} "
                        f"skipped_low_valid={video_skipped_low_valid}"
                    )
                continue

            label = str(segment["label"])
            sample_name = str(segment["sample_name"])
            clip_json = build_clip_json(
                annotations=annotations,
                track_id=track_id,
                start_frame=start_frame,
                end_frame=end_frame,
                sample_name=sample_name,
                category_id=label_to_id[label],
                use_smooth=args.use_smooth,
                joint_score_source=args.joint_score_source,
            )

            out_path = segment_clip_json_path(args.output_root, video_rel_stem, segment_stem)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            with out_path.open("w", encoding="utf-8") as f:
                json.dump(clip_json, f, indent=2, ensure_ascii=False)

            sample_meta = {
                "category_id": label_to_id[label],
                "label": label,
                "source_video": video_rel_stem,
                "source_video_path": str(source_video_path),
                "segment_stem": segment_stem,
                "relative_json_path": str(out_path.relative_to(args.output_root)),
                "kpts_source": kpts_source,
                "track_id": track_id,
                "start_frame": source_start_frame,
                "end_frame": source_end_frame,
                "num_frames": clip_frames,
                "valid_frames": int(valid_stats["valid_frames"]),
                "valid_ratio": float(valid_stats["valid_ratio"]),
                "mean_hand_score": float(valid_stats["mean_hand_score"]),
                "row_idx": int(segment["row_idx"]),
                "annotation_split": None if segment.get("annotation_split") is None else str(segment.get("annotation_split")),
            }
            video_clip_meta[sample_name] = sample_meta
            clip_meta[sample_name] = sample_meta
            video_written += 1
            if should_report_progress(current, total_video_segments):
                log_status(
                    f"{log_prefix} Clip progress {format_progress(current, total_video_segments)} | "
                    f"written={video_written} skipped_short={video_skipped_short} "
                    f"skipped_low_valid={video_skipped_low_valid}"
                )

        video_manifest = {
            "source_video_rel_stem": video_rel_stem,
            "source_video_path": str(source_video_path),
            "signature": current_signature,
            "crop_first": bool(args.crop_first),
            "num_segments_input": len(video_segments),
            "num_samples_written": video_written,
            "num_segments_skipped_short": video_skipped_short,
            "num_segments_skipped_low_valid": video_skipped_low_valid,
            "num_segments_cropped": int(video_crop_stats["num_segments_cropped"]),
            "num_segments_crop_skipped": int(video_crop_stats["num_segments_crop_skipped"]),
            "clips": video_clip_meta,
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        }
        manifest_path.parent.mkdir(parents=True, exist_ok=True)
        with manifest_path.open("w", encoding="utf-8") as f:
            json.dump(video_manifest, f, indent=2, ensure_ascii=False)

        written += video_written
        skipped_short += video_skipped_short
        skipped_low_valid += video_skipped_low_valid
        log_status(
            f"{log_prefix} Finished | written={video_written} skipped_short={video_skipped_short} "
            f"skipped_low_valid={video_skipped_low_valid}"
        )

    category_payload = {
        "categories": categories,
        "annotations": clip_meta,
    }
    with (args.output_root / "category_annotation.json").open("w", encoding="utf-8") as f:
        json.dump(category_payload, f, indent=2, ensure_ascii=False)

    summary = {
        "video_root": str(args.video_root),
        "annotation_file": str(args.annotation_file),
        "kpts_root": str(kpts_root),
        "crop_first": args.crop_first,
        "crop_root": str(crop_root) if args.crop_first else None,
        "output_root": str(args.output_root),
        "categories": categories,
        "num_source_videos": len(segments_by_video),
        "num_source_videos_processed": processed_videos,
        "num_source_videos_skipped": skipped_videos,
        "num_source_videos_rebuilt": rebuilt_videos,
        "num_segments_input": len(segments),
        "num_segments_written": written,
        "num_segments_skipped_short": skipped_short,
        "num_segments_skipped_low_valid": skipped_low_valid,
        "num_segments_cropped": crop_written,
        "num_segments_crop_skipped": crop_skipped,
        "num_annotation_rows_missing_video": len(missing_video_rows),
        "missing_video_rows": missing_video_rows,
        "num_annotation_rows_invalid_frame": len(invalid_frame_rows),
        "invalid_frame_rows": invalid_frame_rows,
        "annotation_format": args.annotation_format,
        "use_smooth": args.use_smooth,
        "joint_score_source": args.joint_score_source,
        "handedness": args.handedness,
        "track_id": args.track_id,
        "min_valid_frames": args.min_valid_frames,
        "min_valid_ratio": args.min_valid_ratio,
        "skip_existing_videos": args.skip_existing_videos,
    }
    with (args.output_root / "prepare_summary.json").open("w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)

    log_status(f"[done] Saved dataset root: {args.output_root}")
    log_status(f"[done] Saved category annotation: {args.output_root / 'category_annotation.json'}")
    log_status(f"[done] Saved prepare summary: {args.output_root / 'prepare_summary.json'}")
    log_status(f"[done] Source videos processed: {processed_videos}")
    log_status(f"[done] Source videos skipped: {skipped_videos}")
    log_status(f"[done] Source videos rebuilt: {rebuilt_videos}")
    log_status(f"[done] Segments written: {written}")
    log_status(f"[done] Segments skipped (too short): {skipped_short}")
    log_status(f"[done] Segments skipped (low valid coverage): {skipped_low_valid}")
    log_status(f"[done] Annotation rows skipped (video not found): {len(missing_video_rows)}")
    log_status(f"[done] Annotation rows skipped (invalid frame range): {len(invalid_frame_rows)}")
    log_status(f"[done] Categories: {categories}")


if __name__ == "__main__":
    main()
